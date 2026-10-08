#!/usr/bin/env python3
"""Taslak dosyasındaki `durum: onay` blokları ornekler.md'ye (eğitim seti) ve cevaplanan.json'a işler.
`durum: red` olanlar sadece cevaplanan.json'a girer (tekrar üretilmesin).
Kullanım: python3 agent/scripts/onayla.py agent/cevaplar/taslak-2026-10-08.md"""
import datetime, json, os, re, sys

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DONE = os.path.join(REPO, "agent", "cevaplar", "cevaplanan.json")
ORNEK = os.path.join(REPO, "agent", "ornekler.md")
BLOCK = re.compile(r"### (?P<id>\w+) · (?P<yazar>.+?) · (?P<puan>\d)★ · (?P<tarih>.*?)\n> (?P<yorum>.*?)\n(?:dil: (?P<dil>.*?)\n)?cevap: ?(?P<cevap>.*?)\n(?:duygu: (?P<duygu>.*?)\n)?(?:jev: (?P<jev>.*?)\n)?durum: (?P<durum>\w+)", re.S)
EXCEL = os.path.join(REPO, "agent", "cevaplar", "yorum-cevaplari.xlsx")
KOLONLAR = [("onay_zamani", "Onay zamanı"), ("taslak_tarihi", "Yorum tarihi"), ("yazar", "Yazar"), ("puan", "Puan"), ("dil", "Dil"), ("yorum", "Yorum"), ("cevap", "Cevap"), ("duygu", "Duygu"), ("jev", "Jev"), ("durum", "Durum"), ("google_review_id", "Google yorum ID"), ("googlea_yazildi", "Google'a yazıldı (elle doldur)")]


def parse(text):
    return [m.groupdict() for m in BLOCK.finditer(text)]


def excel_yaz(done):
    """cevaplanan.json tek doğruluk kaynağı; Excel her onayda baştan üretilir. Son kolon elle doldurulur."""
    from openpyxl import Workbook, load_workbook
    elle = {}
    if os.path.exists(EXCEL):  # elle doldurulan son kolonu koru
        ws = load_workbook(EXCEL).active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[10]: elle[row[10]] = row[11]
    wb = Workbook(); ws = wb.active; ws.title = "Cevaplar"
    ws.append([b for _, b in KOLONLAR])
    for d in done:
        d = dict(d); d["googlea_yazildi"] = elle.get(d.get("google_review_id")) or d.get("googlea_yazildi", "")
        ws.append([("" if d.get(k) is None else d.get(k)) for k, _ in KOLONLAR])
    for col, w in zip("ABCDEFGHIJKL", (19, 12, 22, 5, 5, 60, 70, 9, 14, 7, 30, 16)):
        ws.column_dimensions[col].width = w
    wb.save(EXCEL)


def main():
    if "--test" in sys.argv:
        t = "### abc123 · Ali · 5★ · 2 gün önce\n> süper\ndil: tr\ncevap: sağ olun — Gala Kokoreç Ekibi\nduygu: positive\ndurum: onay\n"
        b = parse(t); assert b and b[0]["id"] == "abc123" and b[0]["durum"] == "onay" and b[0]["dil"] == "tr", b
        print("ok"); return
    blocks = parse(open(sys.argv[1], encoding="utf-8").read())
    done = json.load(open(DONE)) if os.path.exists(DONE) else []
    seen = {d["id"] for d in done}
    n_onay = n_red = 0
    with open(ORNEK, "a", encoding="utf-8") as f:
        for b in blocks:
            if b["durum"] not in ("onay", "red") or b["id"] in seen:
                continue
            done.append({"id": b["id"], "yazar": b["yazar"], "puan": int(b["puan"]), "durum": b["durum"], "cevap": b["cevap"] if b["durum"] == "onay" else None,
                         "yorum": b["yorum"].replace("\n> ", " "), "dil": b.get("dil") or "", "duygu": b.get("duygu") or "", "jev": b.get("jev") or "",
                         "taslak_tarihi": b["tarih"].replace("tarih: ", ""), "onay_zamani": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
                         "google_review_id": b["id"] if b["id"].startswith("Ci9") else ""})
            if b["durum"] == "onay":
                f.write(f"\n## {b['puan']}★\n> {b['yorum']}\ncevap: {b['cevap']}\n"); n_onay += 1
            else:
                n_red += 1
    json.dump(done, open(DONE, "w"), ensure_ascii=False, indent=1)
    excel_yaz(done)
    print(f"onay: {n_onay}, red: {n_red} → ornekler.md, cevaplanan.json ve yorum-cevaplari.xlsx güncellendi")


if __name__ == "__main__":
    main()
